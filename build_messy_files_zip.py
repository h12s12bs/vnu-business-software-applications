# -*- coding: utf-8 -*-
"""
Generate 'Week02_企業混亂命名與版本災難實例包.zip'
Contains 8 realistic messy corporate files that demonstrate version chaos,
allowing students to practice batch audit, renaming, and 4-tier folder organization with Agentic AI.
"""
import os
import zipfile
import docx
import openpyxl
from pptx import Presentation

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DOWNLOADS_DIR = os.path.join(BASE_DIR, "downloads")
TEMP_DIR = os.path.join(BASE_DIR, "temp_messy_files")

os.makedirs(DOWNLOADS_DIR, exist_ok=True)
os.makedirs(TEMP_DIR, exist_ok=True)

# 1. docx files
def create_mock_docx(filepath, title, content):
    doc = docx.Document()
    doc.add_heading(title, 0)
    for p in content:
        doc.add_paragraph(p)
    doc.save(filepath)

# 2. xlsx files
def create_mock_xlsx(filepath, sheet_name, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(filepath)

# 3. pptx files
def create_mock_pptx(filepath, title, subtitle):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    prs.save(filepath)

# 4. txt files
def create_mock_txt(filepath, text):
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(text)

files_to_pack = []

# File 1
f1 = os.path.join(TEMP_DIR, "會議記錄_最新版_final.docx")
create_mock_docx(f1, "門市數位轉型例會初稿", [
    "會議日期：2026年9月中旬（未記清楚確切日期）",
    "紀錄人：前任實習生（已離職）",
    "討論內容：大家討論了全台門市導入智慧零售的事情，營運部說鮮食報廢率有下降，資訊部說要買新主機。"
])
files_to_pack.append(f1)

# File 2
f2 = os.path.join(TEMP_DIR, "會議記錄_真的最終版_主任改過.docx")
create_mock_docx(f2, "門市數位轉型例會修正版", [
    "會議日期：2026年9月21日",
    "主任備註：補充了18家門市預算，每台85,000元，請財務部核對。",
    "待辦事項：Susan負責教育訓練，Kevin負責合約。"
])
files_to_pack.append(f2)

# File 3
f3 = os.path.join(TEMP_DIR, "2026年9月門市預算表_小美修改版_最終不改版.xlsx")
create_mock_xlsx(f3, "門市預算", 
    ["門市分區", "門市數", "設備單價", "總金額(元)", "審核狀態"],
    [
        ["北區門市", 6, 85000, 510000, "已審核"],
        ["中區門市", 6, 85000, 510000, "待覆核"],
        ["南區門市", 6, 85000, 510000, "待覆核"],
        ["全台合計", 18, 85000, 1530000, "通過"]
    ]
)
files_to_pack.append(f3)

# File 4
f4 = os.path.join(TEMP_DIR, "董事會報告_David總經理最終確定版(1).pptx")
create_mock_pptx(f4, "全通路智慧零售轉型三年旗艦案", "萬能企業集團 董事會提報 ｜ 總經理室")
files_to_pack.append(f4)

# File 5
f5 = os.path.join(TEMP_DIR, "智慧零售企劃草案_v2_new_final_FINAL.docx")
create_mock_docx(f5, "智慧零售轉型三年旗艦企劃案（草案）", [
    "本企劃案旨在推動全通路智鏈零售升級工程，預算規模新台幣5,000萬元。",
    "包含鮮食動態折扣系統、物聯網冷鏈溫控及跨門市智慧調度。"
])
files_to_pack.append(f5)

# File 6
f6 = os.path.join(TEMP_DIR, "門市POS合約條款_Linda審核版_最新.docx")
create_mock_docx(f6, "智慧零售系統供應商採購合約要點", [
    "合約編號：CTR-2026-POS-01",
    "SLA條款：系統正常運作時間保證達 99.9%。",
    "違約罰則：若延遲上線，每日處以合約總額千分之二罰款。"
])
files_to_pack.append(f6)

# File 7
f7 = os.path.join(TEMP_DIR, "北中南18家門市名冊_小林手打備份.xlsx")
create_mock_xlsx(f7, "門市清單",
    ["門市代碼", "門市名稱", "負責主管", "聯絡電話", "營運區域"],
    [
        ["ST001", "台北站前旗艦店", "陳店長", "02-2311-0001", "北部"],
        ["ST005", "中壢萬能中原店", "林店長", "03-4511-0005", "北部"],
        ["ST008", "台中逢甲文華店", "黃店長", "04-2451-0008", "中部"],
        ["ST014", "高雄巨蛋裕誠店", "張店長", "07-5551-0014", "南部"]
    ]
)
files_to_pack.append(f7)

# File 8
f8 = os.path.join(TEMP_DIR, "重要通知_千萬不要刪除.txt")
create_mock_txt(f8, "各部門注意：\n所有專案檔案請遵循四層資料夾與 ISO 8601 命名規範，請勿再儲存為『最新版_真的最終版』，否則將依公司資訊治理規章進行稽核！")
files_to_pack.append(f8)

# Instruction file
f_readme = os.path.join(TEMP_DIR, "README_練習說明_請交給AI治理.txt")
create_mock_txt(f_readme, """【萬能科技大學 商業軟體應用 ── 企業混亂檔案治理實戰】

這批檔案模擬真實企業中最常見的「版本混亂災難」：
- 主觀命名（最新版、真的最終版、小美修改版、最終不改版）
- 檔名無日期（完全無法用作業系統排序先後順序）
- 檔名包含 (1)、_FINAL 等隨意修訂標籤
- 所有重要公文、原始數據、未定草案與最終發布檔全部混雜在同一個資料夾！

【學生任務】：
請打開 Antigravity 桌面版（或 Google AI Studio 網頁版），將檔名清單餵給 AI：
「請扮演企業數位資產架構師，依據 ISO 8601 (YYYY-MM-DD) 規範建立前後對照審計表，並自動歸檔至 00_Admin ~ 03_Deliverables 四層樹狀結構！」
""")
files_to_pack.append(f_readme)

# Pack into zip
zip_target = os.path.join(DOWNLOADS_DIR, "Week02_企業混亂命名與版本災難實例包.zip")
with zipfile.ZipFile(zip_target, "w", zipfile.ZIP_DEFLATED) as z:
    for f in files_to_pack:
        arcname = os.path.basename(f)
        z.write(f, arcname)

# Cleanup temp files
for f in files_to_pack:
    if os.path.exists(f):
        os.remove(f)
if os.path.exists(TEMP_DIR):
    os.rmdir(TEMP_DIR)

print(f"Successfully generated: {zip_target} ({os.path.getsize(zip_target):,} bytes)")
